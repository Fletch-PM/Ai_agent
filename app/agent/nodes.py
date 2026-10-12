from pathlib import Path
import re

from docx import Document
from langchain_core.messages import AIMessage
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.output_parsers import StrOutputParser

from app.prompts.agent_identity import agent_identity_sp
from app.prompts.sub_agents.author import author_agent
from app.prompts.sub_agents.researcher import researcher_agent
from app.models.claude import gen_llm
from app.tools.rag.retriever import search as rag_search
from app.tools.websearch.search_intents import extract_search_query, match_search_intent
from app.tools.websearch.web_search import web_search
from app.tools.weather.weather_search import is_weather_query, weather_search

agent_template = ChatPromptTemplate.from_messages(
    [
        ("system", agent_identity_sp),
        (
            "human",
            "Tool status:\n{tool_instruction}\n\nContext:\n{context}\n\nQuestion:\n{user_input}",
        ),
    ]
)
agent_chain = agent_template | gen_llm | StrOutputParser() if gen_llm else None


def _ask(prompt, request, context=""):
    if agent_chain is None:
        return context or "Please set ANTHROPIC_API_KEY before starting the agent."
    return agent_chain.invoke({"user_input": request, "context": context, "tool_instruction": prompt})


def _research(request, reference_context=""):
    query = extract_search_query(request) if match_search_intent(request) else request
    if is_weather_query(query):
        results = weather_search(query)
    else:
        results = web_search(query)
    search_context = "\n\n".join(
        f"Title: {item.get('title', 'Reference')}\nSummary: {item.get('snippet', item.get('parent', item.get('text', '')))}\nURL: {item.get('url', 'local reference')}"
        for item in results
    )
    context = "\n\n".join(filter(None, (reference_context, search_context)))
    return _ask(researcher_agent or "Use the research results to answer the request.", request, context)


def _is_rag_request(request):
    return bool(
        re.search(
            r"\b(?:my\s+(?:profile|resume|(?:professional\s+)?background|experience|skills|name)|about\s+me|who\s+am\s+i|reference)\b",
            request,
            re.I,
        )
    )


def orchestrator_node(state):
    request = state["user_input"]
    context = "\n\n".join(
        filter(
            None,
            (item.get("parent", item.get("text", "")) for item in rag_search(request)),
        )
    )
    lowered_request = request.lower()
    if context and _is_rag_request(request) and "write" not in lowered_request and "search" not in lowered_request:
        answer = _ask(
            "Answer the user's question using only the supplied vector database context. "
            "If it does not contain the answer, say so clearly.",
            request,
            context,
        )
        return {"route": "end", "rag_context": context, "messages": [AIMessage(content=answer)]}
    return {
        "route": "author" if "write" in request.lower() else "researcher",
        "rag_context": context,
    }


def researcher_node(state):
    return {"messages": [AIMessage(content=_research(state["user_input"], state["rag_context"]))]}


def author_node(state):
    request = state["user_input"]
    research = _research(request, state["rag_context"]) if "search" in request.lower() else state["rag_context"]
    paper = _ask(author_agent or "Write the requested paper.", request, research)
    output_dir = Path("documents")
    output_dir.mkdir(exist_ok=True)
    filename = re.sub(r"[^a-z0-9]+", "_", request.lower()).strip("_")[:50] or "paper"
    path = output_dir / f"{filename}.docx"
    document = Document()
    document.add_heading("Research Paper", 0)
    document.add_paragraph(paper)
    document.save(path)
    return {"messages": [AIMessage(content=f"I wrote the paper to `{path}`.")]} 

